"""Download the RAG benchmark corpus and check every hash."""

import argparse
import hashlib
import re
import sys
from pathlib import Path

import httpx2
import yaml

from evals.rag.retrieval import CORPUS_DIR

CORPUS = Path(__file__).with_name("corpus.yaml")


def markdown_to_docx(markdown: str, target: Path) -> None:
    """Headings and paragraphs only; enough to exercise the DOCX parser."""
    from docx import Document

    doc = Document()
    for block in re.split(r"\n\s*\n", markdown):
        text = block.strip()
        if not text or text.startswith("{*"):
            continue
        heading = re.match(r"^(#{1,6})\s+(.*?)(\s*\{.*\})?$", text)
        if heading:
            doc.add_heading(heading.group(2), level=len(heading.group(1)))
        else:
            doc.add_paragraph(" ".join(text.split()))
    doc.save(str(target))


def fetch(directory: Path) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    entries = yaml.safe_load(CORPUS.read_text())["documents"]
    paths: list[Path] = []
    with httpx2.Client(follow_redirects=True, timeout=60) as client:
        for entry in entries:
            target = directory / entry["name"]
            response = client.get(entry["url"])
            response.raise_for_status()
            digest = hashlib.sha256(response.content).hexdigest()
            if digest != entry["sha256"]:
                sys.exit(f"{entry['name']}: sha256 {digest} != {entry['sha256']}")
            if entry.get("convert") == "docx":
                markdown_to_docx(response.text, target)
            else:
                target.write_bytes(response.content)
            print(f"ok  {entry['name']}")
            paths.append(target)
    return paths


def main(argv: list[str] | None = None) -> int:
    argparse.ArgumentParser(description=__doc__).parse_args(argv)
    fetch(CORPUS_DIR)
    return 0


if __name__ == "__main__":
    sys.exit(main())
